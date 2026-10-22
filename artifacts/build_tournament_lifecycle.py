from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION

OUT = "artifacts/TOM Chrispyjams Turnier Lifecycle.docx"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def border(cell, color="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = tcBorders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tcBorders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)

def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.space_before = Pt(3)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    border(cell)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p

def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for idx, header in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.width = Inches(widths[idx])
        shade(cell, "17365D")
        set_cell_text(cell, header, bold=True, color="FFFFFF")
    for row_idx, row in enumerate(rows):
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].width = Inches(widths[idx])
            if row_idx % 2:
                shade(cells[idx], "EEF4FA")
            set_cell_text(cells[idx], value)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10.5)
for style_name in ("Title", "Heading 1", "Heading 2"):
    styles[style_name].font.name = "Aptos Display"
    styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].font.size = Pt(26)
styles["Heading 1"].font.size = Pt(16)
styles["Heading 2"].font.size = Pt(12)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.add_run("Chrispyjams  |  TOM Turnierbetrieb").font.size = Pt(8)

p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("TOM Chrispyjams Turnier Lifecycle")
subtitle = doc.add_paragraph("Ablauf für Veranstalter, Head Judge und Spieler")
subtitle.runs[0].font.size = Pt(13)
subtitle.runs[0].font.color.rgb = RGBColor(68, 68, 68)
body(doc, "Dieses Dokument beschreibt den operativen Ablauf eines Pokémon TCG Turniers mit TOM als offizieller Pairing Quelle und Chrispyjams als Spieleroberfläche. TOM bleibt für Turnierstruktur und die Erstellung der nächsten Runde maßgeblich. Chrispyjams zeigt Standings, verteilt Pairings, sammelt bestätigte Ergebnisse und liefert die abgeschlossene Pod TDF zurück an TOM.")

heading(doc, "Rollen und Systeme")
table(doc, ["Bereich", "System", "Verantwortung"], [
    ("Turnierstruktur", "TOM", "Teilnehmer, Pods, Swiss und Top Cut anlegen; Pairings für jede Runde erzeugen."),
    ("Live Betrieb", "Chrispyjams Turniermanager", "TDF importieren, Standings prüfen, Runde freigeben, Ergebnisse überwachen und Pod Export erzeugen."),
    ("Spieler", "Chrispyjams Spieleransicht", "Pairing, Tisch und Gegner sehen; Sieg, Niederlage oder Unentschieden melden."),
    ("Finale Entscheidung", "Head Judge", "Konflikte und Sonderfälle im Turniermanager manuell entscheiden."),
], [1.35, 1.7, 3.95])

heading(doc, "Kurzprinzip")
body(doc, "Jede Runde läuft in einer klaren Schleife: TOM paart, Chrispyjams veröffentlicht und sammelt, TOM erstellt anschließend die nächste Runde. Bei mehreren Altersdivisionen ist jede Division ein eigener Pod. Ein Pod darf exportiert werden, sobald seine aktuelle Runde vollständig entschieden ist, auch wenn ein anderer Pod noch spielt.")

heading(doc, "Vor dem Turnier")
table(doc, ["Schritt", "Wo", "Aktion und Ergebnis"], [
    ("1. Event anlegen", "Chrispyjams Admin Events", "Datum, Ort, Kapazität und Registrierung anlegen. Der Eintrag ist die Chrispyjams Klammer für das Turnier."),
    ("2. Teilnehmer erfassen", "TOM", "Spieler in TOM registrieren oder importieren. Pokémon Spieler ID und Geburtsdatum müssen korrekt sein; sie werden für den Gastzugang verwendet."),
    ("3. Pods festlegen", "TOM", "Junioren, Senioren und Masters nach Bedarf als getrennte Pods anlegen. TOM erzeugt die offizielle Turnierstruktur."),
    ("4. Erste Paarung", "TOM", "Runde 1 paaren und die entstandene TDF Datei speichern."),
    ("5. TDF importieren", "Chrispyjams Eventverwaltung → Turnier starten", "Am richtigen Event den Live-Workspace öffnen, TDF hochladen und Pod Standings sowie Tische kurz prüfen."),
    ("6. Turnier freigeben", "Chrispyjams Turniermanager", "Erst diese Aktion öffnet die persönliche Turnieransicht, das aktuelle Standing und den globalen QR-Zugang. Anschließend QR Display auf dem Fernseher öffnen."),
], [1.25, 1.7, 4.05])

heading(doc, "Ablauf pro Runde")
table(doc, ["Reihenfolge", "System und Maske", "Was passiert"], [
    ("1", "TOM", "Judge beendet die vorherige Runde in TOM und erzeugt die nächste Paarung für jeden aktiven Pod."),
    ("2", "Chrispyjams Turniermanager", "Neue TDF importieren. Die Übersicht zeigt Standings; jeder Rundentab zeigt Pod, Tisch, Gegner und Ergebnisstatus."),
    ("3", "Chrispyjams Turniermanager", "Im passenden Pod Runde starten. Erst dann sehen Spieler die neuen Pairings in ihrer Turnieransicht; Pods können unabhängig freigegeben werden."),
    ("4", "Fernseher QR Display", "Spieler scannen den Code. Nicht registrierte Chrispyjams Nutzer geben Pokémon Spieler ID und Geburtsjahr aus der TOM Datei ein."),
    ("5", "Chrispyjams Spieleransicht", "Spieler sehen Tisch, Gegner und Runde. Nach dem Match melden beide unabhängig Sieg, Niederlage oder Unentschieden."),
    ("6", "Chrispyjams Turniermanager", "Übereinstimmende Meldungen werden übernommen. Abweichende Meldungen erscheinen als Konflikt und werden durch den Judge entschieden."),
    ("7", "Chrispyjams Turniermanager", "Wenn alle Matches eines Pods entschieden sind, Pod TDF exportieren."),
    ("8", "TOM", "Exportierte Pod TDF importieren beziehungsweise Ergebnisse übernehmen. TOM berechnet die nächste Paarung. Danach beginnt die Schleife erneut."),
], [0.9, 2.0, 4.1])

heading(doc, "Ergebnisregeln")
body(doc, "Spielereingaben sind keine einseitige Bestätigung. Chrispyjams trägt ein Ergebnis nur ein, wenn beide Spieler dieselbe Spielperspektive melden: Sieg gegen Niederlage oder zweimal Unentschieden. Bei Abweichung bleibt der Status offen und der Head Judge entscheidet im Rundentab. Der Judge kann außerdem jederzeit ein Ergebnis direkt setzen, etwa bei Double Game Loss, Drop oder einer Schiedsrichterentscheidung.")

heading(doc, "Mehrere Pods")
body(doc, "Ein Cup kann zeitgleich mehrere Pods enthalten. In der Übersicht erscheint für jeden Pod eine eigene Standing Tabelle und ein eigener Exportknopf. Der Export enthält nur diesen Pod und die zugehörigen TOM Spieler. So kann beispielsweise Juniors Seniors bereits in TOM weitergepaart werden, während Masters noch spielt.")

heading(doc, "Abschluss")
table(doc, ["Schritt", "Wo", "Ergebnis"], [
    ("Letzte Runde abschließen", "Chrispyjams", "Alle offenen Ergebnisse pro Pod prüfen und gegebenenfalls Konflikte lösen."),
    ("Finalen Stand exportieren", "Chrispyjams", "Für jeden Pod die fertige TDF exportieren und in TOM übernehmen."),
    ("Turnier schließen", "TOM und Chrispyjams Admin Events", "TOM Ergebnisse finalisieren; Event in Chrispyjams als abgeschlossen markieren. Spielerhistorie bleibt sichtbar."),
], [1.55, 1.9, 3.55])

heading(doc, "Betriebshinweise")
body(doc, "Vor jeder Freigabe nur kurz prüfen, ob die hochgeladene TDF zum richtigen Event und zur erwarteten Runde gehört. Bei Änderungen in TOM immer die neue TDF erneut importieren; TOM ist die Quelle für Pairings. Chrispyjams ist die Quelle für während der Runde bestätigte Spielerergebnisse bis zum jeweiligen Pod Export.")

for s in doc.sections:
    footer = s.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Interne Betriebsanleitung  |  Stand Oktober 2026").font.size = Pt(8)

doc.save(OUT)
